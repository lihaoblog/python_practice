"""
    该功能是把word的表格里面的数据写到excel表格里面，注意是必须要word里面用了表格
"""
from docx import Document
import pandas as pd

document=Document(r'C:\Users\lihao\Desktop\dify\测试表格文档.docx')

tables=document.tables

# excel=tables[0].cell(0,1).text
# print(excel)

# for i in range(0,len(tables)):
#     chengshi=tables[i].cell(0,0).text
#     mingcheng=tables[i].cell(0,1).text
#     mianji = tables[i].cell(0, 2).text
#     print(chengshi,mingcheng,mianji)

# 创建空的列表来存储数据，我们把每个切分的表，取同类型的数据放入一个表格中来
chengshi_list=[]
mingcheng_list=[]
mianji_list=[]

# 用append来将多个切分的表中的同类型数据放入一个人列表中，后续写入
for i in range(0,len(tables)):
    chengshi_list.append(tables[i].cell(1,0).text)
    mingcheng_list.append(tables[i].cell(1,1).text)
    mianji_list.append(tables[i].cell(1, 2).text)
print(chengshi_list,mingcheng_list,mianji_list)

info_dict={
    '城市':chengshi_list,
    '名称':mingcheng_list,
    '面积':mianji_list
}
# DataFrame关键字可以把字段转成想要的格式
df=pd.DataFrame(info_dict)
print(df)
df.to_excel(r'C:\Users\lihao\Desktop\dify\word转excel.xlsx',index=False)
print('success')